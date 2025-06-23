import os
import re
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

def is_hindi(text):
    # Detect Devanagari script (Hindi)
    return bool(re.search(r'[\u0900-\u097F]', text))

def writeContentToDoc(file_path, content):
    """
    Writes or appends English and Hindi content to a .docx file,
    using a font that supports Devanagari script.
    """
    # Ensure directory exists
    folder = os.path.dirname(file_path)
    if folder and not os.path.exists(folder):
        os.makedirs(folder)

    # Open or create document
    if os.path.exists(file_path):
        doc = Document(file_path)
        doc.add_page_break()
    else:
        doc = Document()

    # Normalize content
    if isinstance(content, str):
        content = [content]

    for para in content:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(para)

        # Use Nirmala UI which supports both Hindi and English
        run.font.name = 'Nirmala UI'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Nirmala UI')
        run.font.size = Pt(14)

    doc.save(file_path)
    # print(f"Written to: {file_path}")
